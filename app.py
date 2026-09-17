
import json
import os
import random
import re
import zipfile
from pathlib import Path

from flask import Flask, jsonify, render_template, request, send_file
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
import arabic_reshaper
from bidi.algorithm import get_display

BASE_DIR = Path(__file__).resolve().parent
CURRICULUM_DATA = BASE_DIR / "data" / "curriculum.json"
QUESTIONS_DATA = BASE_DIR / "data" / "questions.json"
app = Flask(__name__)

TYPE_LABELS = {
    "mcq": "اختيار من متعدد",
    "tf": "صح أو خطأ",
    "fill": "أكمل الفراغ",
    "match": "توصيل الكلمة بالمصطلح المناسب",
}
DIFFICULTY_LABELS = {"easy": "سهل", "medium": "متوسط", "hard": "صعب"}


def load_json(path, default):
    if not path.exists():
        return default
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def curriculum():
    return load_json(CURRICULUM_DATA, {})


def question_bank():
    data = load_json(QUESTIONS_DATA, [])
    return data.get("questions", []) if isinstance(data, dict) else data


def int_count(value):
    try:
        return max(0, min(100, int(value or 0)))
    except Exception:
        return 0


def strip_option_prefix(value):
    value = str(value or "").strip()
    patterns = [
        r"^\s*[1-4١-٤]\s*[\.\-\)\:]\s*",
        r"^\s*[أبجد]\s*[\.\-\)\:]\s*",
    ]
    for pattern in patterns:
        value = re.sub(pattern, "", value, count=1)
    return value.strip()


def numbered_mcq_options(options):
    clean = [strip_option_prefix(x) for x in (options or [])]
    clean = [x for x in clean if x][:4]
    return [f"{i}. {value}" for i, value in enumerate(clean, start=1)]


def compact_term(term):
    term = str(term or "").strip()
    return term.replace("الفصل الدراسي ", "").strip() or term


def normalize_types(payload):
    values = payload.get("types") or []
    return [x for x in values if x in TYPE_LABELS]


def normalize_lessons(payload):
    return [str(x).strip() for x in (payload.get("lessons") or []) if str(x).strip()]


def matches_value(item_value, requested_value):
    return not item_value or str(item_value).strip() == str(requested_value or "").strip()


def eligible_questions(payload, difficulty):
    lessons = set(normalize_lessons(payload))
    types = set(normalize_types(payload))
    selected = []
    for q in question_bank():
        if not isinstance(q, dict):
            continue
        if q.get("difficulty") != difficulty:
            continue
        if q.get("type") not in types:
            continue
        if q.get("lesson") not in lessons:
            continue
        if not matches_value(q.get("stage"), payload.get("stage")):
            continue
        if not matches_value(q.get("grade"), payload.get("grade")):
            continue
        if not matches_value(q.get("track"), payload.get("track")):
            continue
        if not matches_value(q.get("term"), payload.get("term")):
            continue
        if not matches_value(q.get("subject"), payload.get("subject")):
            continue
        selected.append(q)
    return selected


def balanced_pick(pool, count, lessons, used_ids):
    if count <= 0:
        return []
    by_lesson = {lesson: [] for lesson in lessons}
    for q in pool:
        qid = str(q.get("id") or "")
        if qid and qid in used_ids:
            continue
        lesson = q.get("lesson")
        if lesson in by_lesson:
            by_lesson[lesson].append(q)
    for values in by_lesson.values():
        random.shuffle(values)

    result = []
    active = list(lessons)
    random.shuffle(active)
    while len(result) < count:
        progressed = False
        for lesson in active:
            if len(result) >= count:
                break
            if by_lesson.get(lesson):
                q = by_lesson[lesson].pop()
                result.append(q)
                qid = str(q.get("id") or "")
                if qid:
                    used_ids.add(qid)
                progressed = True
        if not progressed:
            break
    return result


def generate_from_bank(payload):
    lessons = normalize_lessons(payload)
    types = normalize_types(payload)
    difficulties = payload.get("difficulties") or {}

    if not lessons:
        raise ValueError("اختر درساً واحداً على الأقل.")
    if not types:
        raise ValueError("اختر نوعاً واحداً على الأقل من الأسئلة.")

    requested = {key: int_count(difficulties.get(key, 0)) for key in DIFFICULTY_LABELS}
    total = sum(requested.values())
    if total <= 0:
        raise ValueError("حدد عدد الأسئلة السهلة أو المتوسطة أو الصعبة.")
    if total > 100:
        raise ValueError("الحد الأقصى للاختبار الواحد هو 100 سؤال.")

    picked = []
    used_ids = set()
    shortages = []
    for difficulty, count in requested.items():
        pool = eligible_questions(payload, difficulty)
        chosen = balanced_pick(pool, count, lessons, used_ids)
        picked.extend(chosen)
        if len(chosen) < count:
            shortages.append(
                f"{DIFFICULTY_LABELS[difficulty]}: المطلوب {count} والمتاح {len(chosen)}"
            )

    if shortages:
        raise ValueError(
            "بنك الأسئلة لا يحتوي عدداً كافياً للاختيارات الحالية. "
            + " | ".join(shortages)
        )

    random.shuffle(picked)
    questions = []
    for i, q in enumerate(picked, start=1):
        qtype = q.get("type", "")
        options = q.get("options") or []
        if qtype == "mcq":
            options = [strip_option_prefix(x) for x in options][:4]
        questions.append({
            "n": i,
            "id": q.get("id", ""),
            "type": TYPE_LABELS.get(qtype, qtype),
            "typeKey": qtype,
            "difficulty": DIFFICULTY_LABELS.get(q.get("difficulty"), q.get("difficulty", "")),
            "lesson": q.get("lesson", ""),
            "text": q.get("text", ""),
            "options": options,
            "answer": q.get("answer", ""),
            "explanation": q.get("explanation", ""),
        })
    return questions


@app.get("/")
def home():
    return render_template("index.html", data=curriculum())


@app.get("/api/health")
def health():
    return jsonify({"ok": True, "mode": "question_bank", "questions": len(question_bank())})


@app.post("/api/generate")
def gen():
    try:
        payload = request.get_json(force=True) or {}
        questions = generate_from_bank(payload)
        return jsonify({"ok": True, "questions": questions, "message": "تم إنشاء الاختبار من بنك الأسئلة."})
    except Exception as e:
        return jsonify({"ok": False, "error": str(e)}), 400


def set_rtl(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    pPr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    pPr.append(bidi)


def add_docx_line(doc, text="", bold=False, size=12, center=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.RIGHT
    if not center:
        set_rtl(p)
    r = p.add_run(str(text))
    r.bold = bold
    r.font.name = "Arial"
    r.font.size = Pt(size)
    return p


def docx_bytes(qs, meta):
    d = Document()
    d.styles["Normal"].font.name = "Arial"
    d.styles["Normal"].font.size = Pt(12)
    title = meta.get("title") or meta.get("documentType") or "مستند تعليمي"
    add_docx_line(d, title, bold=True, size=18, center=True)
    info = (
        f"المادة: {meta.get('subject') or '—'}  |  الصف: {meta.get('grade') or '—'}  |  "
        f"الفصل الدراسي: {compact_term(meta.get('term')) or '—'}  |  المعلم: {meta.get('teacher') or '—'}"
    )
    add_docx_line(d, info, bold=True, size=11)
    add_docx_line(d, "اسم الطالب/ـة: ______________________________", bold=True)
    d.add_paragraph("")
    for q in qs:
        add_docx_line(d, f"{q.get('n', '')}. {q.get('text', '')}", bold=True)
        opts = q.get("options") or []
        if opts:
            if q.get("type") == "اختيار من متعدد" or q.get("typeKey") == "mcq":
                add_docx_line(d, "     ".join(numbered_mcq_options(opts)), size=11)
            else:
                for x in opts:
                    add_docx_line(d, x, size=11)
        d.add_paragraph("")
    b = io.BytesIO()
    d.save(b)
    b.seek(0)
    return b.getvalue()


def find_arabic_font():
    candidates = [
        BASE_DIR / "static" / "fonts" / "DejaVuSans.ttf",
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        Path("/usr/share/fonts/truetype/noto/NotoNaskhArabic-Regular.ttf"),
        Path("/usr/share/fonts/opentype/noto/NotoNaskhArabic-Regular.ttf"),
    ]
    for p in candidates:
        if p.exists():
            return str(p)
    return None


ARABIC_FONT_PATH = find_arabic_font()
PDF_FONT = "Helvetica"
if ARABIC_FONT_PATH:
    try:
        pdfmetrics.registerFont(TTFont("ArabicFont", ARABIC_FONT_PATH))
        PDF_FONT = "ArabicFont"
    except Exception:
        pass


def ar(text):
    text = str(text or "")
    if PDF_FONT == "Helvetica":
        return text
    return get_display(arabic_reshaper.reshape(text))


def wrap_text(text, max_chars=85):
    words = str(text or "").split()
    lines, current, length = [], [], 0
    for word in words:
        extra = len(word) + (1 if current else 0)
        if length + extra > max_chars and current:
            lines.append(" ".join(current))
            current, length = [word], len(word)
        else:
            current.append(word)
            length += extra
    if current:
        lines.append(" ".join(current))
    return lines or [""]


def pdf_bytes(qs, meta):
    b = io.BytesIO()
    c = canvas.Canvas(b, pagesize=A4)
    w, h = A4
    y = h - 45

    def new_page():
        nonlocal y
        c.showPage()
        y = h - 45

    def draw_right(text, font_size=11, gap=16):
        nonlocal y
        c.setFont(PDF_FONT, font_size)
        for line in wrap_text(text):
            if y < 50:
                new_page()
                c.setFont(PDF_FONT, font_size)
            c.drawRightString(w - 40, y, ar(line))
            y -= gap

    title = meta.get("title") or meta.get("documentType") or "مستند تعليمي"
    draw_right(title, 17, 22)
    info = (
        f"المادة: {meta.get('subject') or '—'} | الصف: {meta.get('grade') or '—'} | "
        f"الفصل الدراسي: {compact_term(meta.get('term')) or '—'} | المعلم: {meta.get('teacher') or '—'}"
    )
    draw_right(info, 10, 17)
    draw_right("اسم الطالب/ـة: ______________________________", 11, 19)
    y -= 8
    for q in qs:
        draw_right(f"{q.get('n', '')}. {q.get('text', '')}", 11, 17)
        opts = q.get("options") or []
        if opts:
            if q.get("type") == "اختيار من متعدد" or q.get("typeKey") == "mcq":
                draw_right("     ".join(numbered_mcq_options(opts)), 9, 15)
            else:
                for x in opts:
                    draw_right(x, 9, 14)
        y -= 7
    c.save()
    b.seek(0)
    return b.getvalue()


@app.post("/api/export")
def export():
    try:
        p = request.get_json(force=True) or {}
        qs, meta = p.get("questions") or [], p.get("meta") or {}
        fmt = (p.get("format") or "pdf").lower()
        if fmt == "pdf":
            return send_file(io.BytesIO(pdf_bytes(qs, meta)), as_attachment=True, download_name="المستند.pdf", mimetype="application/pdf")
        if fmt == "docx":
            return send_file(io.BytesIO(docx_bytes(qs, meta)), as_attachment=True, download_name="المستند.docx", mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        z = io.BytesIO()
        with zipfile.ZipFile(z, "w", zipfile.ZIP_DEFLATED) as f:
            f.writestr("المستند.pdf", pdf_bytes(qs, meta))
            f.writestr("المستند.docx", docx_bytes(qs, meta))
        z.seek(0)
        return send_file(z, as_attachment=True, download_name="المستندات.zip", mimetype="application/zip")
    except Exception as e:
        return jsonify({"ok": False, "error": str(e)}), 400


if __name__ == "__main__":
    port = int(os.environ.get("PORT", "5000"))
    app.run(host="0.0.0.0", port=port)
